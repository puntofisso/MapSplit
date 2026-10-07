"""Build the review workbook (review/MapSplit-review.xlsx) from the pipeline's
outputs, for the owner to mark up — uploaded to Google Drive as a Sheet.

    python3 tools/pipeline/review_workbook.py

Reads data/stats.json, review/questions.csv (run measure.mjs first),
schedule/themes.json, review/themes-normal.csv and data/games/2027.json.
Proposed wording here is a DRAFT for the owner to accept or change; nothing
in the game reads this file.
"""
import csv
import datetime
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
P = lambda *a: os.path.join(ROOT, *a)
YEAR = 2027
EPOCH = datetime.date(2026, 8, 25)            # GAME.epoch: puzzle #1
PREVIEW = 'http://localhost:8765/MapSplit/Play/?date='

cat = json.load(open(P('data/stats.json')))
q = {(r['statistic'], r['axis']): r for r in csv.DictReader(open(P('review/questions.csv')))}
cfg = json.load(open(P('schedule/themes.json')))
tcount = {r['theme']: r for r in csv.DictReader(open(P('review/themes-normal.csv')))}
games = json.load(open(P('data/games/%d.json' % YEAR)))

# Draft wording for the questions long enough to wrap on a phone.
SHORTER = {
    'repd-battery-pipeline': 'planned battery storage capacity',
    'census-home-commercial': 'households living above a shop or office',
    'census-sick': 'people too ill or disabled to work',
    'census-home-converted': 'households in bedsits and converted houses',
    'census-qual-apprentice': 'people qualified through an apprenticeship',
    'repd-solar-pipeline': 'planned solar farm capacity',
    'census-carers-50h': 'people giving 50+ hours of unpaid care a week',
    'census-qual-degree': 'people with a degree or higher',
    'census-commute-metro': 'people who commute by tube, metro or tram',
}
# Mass nouns take "is" in the takeaway; everything else is a plural count.
SINGULAR = {'pop-all', 'census-population', 'hmrc-tax-amount', 'repd-wind-onshore',
            'repd-wind-offshore', 'repd-solar', 'repd-solar-pipeline', 'repd-battery',
            'repd-battery-pipeline'}

days_per_stat, themes_of = {}, {}
for d in games['days']:
    for sid, _ in d['rounds']:
        days_per_stat[sid] = days_per_stat.get(sid, 0) + 1
for t in cfg['themes']:
    for sid in t['stats']:
        themes_of.setdefault(sid, []).append(t['title'])

HEAD = Font(bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor='2F6F9F')
EDIT_FILL = PatternFill('solid', fgColor='FFF7E6')     # the columns the owner fills in
WRAP = Alignment(wrap_text=True, vertical='top')


def sheet(wb, title, header, rows, widths, edit_cols=(), keep_col=None, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(header)
    for r in rows:
        ws.append(r)
    for i, c in enumerate(ws[1], 1):
        c.font, c.fill, c.alignment = HEAD, HEAD_FILL, WRAP
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
            if c.column in edit_cols:
                c.fill = EDIT_FILL
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions
    if keep_col:
        dv = DataValidation(type='list', formula1='"Keep,Change,Drop"', allow_blank=True)
        ws.add_data_validation(dv)
        col = ws.cell(1, keep_col).column_letter
        dv.add('%s2:%s%d' % (col, col, ws.max_row))
    return ws


wb = Workbook()

readme = [
    ['MapSplit — review workbook', ''],
    ['Generated', datetime.date.today().isoformat() + ' from data/stats.json, the measurements in review/ and data/games/%d.json.' % YEAR],
    ['How to use it', 'Fill in the shaded (cream) columns only: Keep / Change / Drop, your wording, notes. Claude reads this Sheet back and applies the changes to sources/*/source.json and schedule/themes.json, then rebuilds and regenerates the schedule.'],
    ['Statistics', 'One row per statistic (84). "Proposed question" is a DRAFT only where the current one is long enough to wrap on a phone. "Proposed takeaway" is how the after-guess fact would read with the question wording instead of the label (fixes units such as "(MW)" appearing in sentences). Lines are the 50/50 answer; "km vs population" is how far that sits from the population\'s own line — the bigger, the more surprising.'],
    ['Themes', 'One row per theme. Valid days = distinct (three statistics, axis) days in Normal mode with all lines at least 18 km apart.'],
    ['Calendar %d' % YEAR, 'The generated year. The preview link opens that day in the game on your machine (run: php -S localhost:8765 -t .. from the MapSplit folder). Preview scores are kept separate from real ones.'],
    ['Not here', 'The rest of 2026 is regenerated from the launch date, so it is not listed.'],
]
ws = sheet(wb, 'Read me', ['', ''], readme[0:0], [22, 110], first=True)
for r in readme:
    ws.append(r)
ws.delete_rows(1)
for row in ws.iter_rows():
    row[0].font = Font(bold=True)
    for c in row:
        c.alignment = WRAP
ws.freeze_panes = None

srows = []
for sid in sorted(cat, key=lambda s: (cat[s]['family'], s)):
    e = cat[sid]
    ns, we = q[(sid, 'ns')], q[(sid, 'we')]
    question = e['question']
    proposed = SHORTER.get(sid, '')
    noun = proposed or question
    verb = 'is' if sid in SINGULAR else 'are'
    take = 'Half of %s %s %s of %s — a line through %s.' % (
        noun, verb, 'north', ns['line_at'], ns['place'] or '…')
    srows.append([
        sid, e['label'], e['family'], question, proposed, take,
        ns['line_at'], ('%s km %s' % (ns['km_from_population'], ns['toward'].lower())) if ns['toward'] else '—',
        we['line_at'], ('%s km %s' % (we['km_from_population'], we['toward'].lower())) if we['toward'] else '—',
        '; '.join(themes_of.get(sid, [])) or '(in no theme)', days_per_stat.get(sid, 0),
        e['points'], e['source'], e['derivedLicence']['name'], ' | '.join(e['caveats']),
        '', '', '',
    ])
sheet(wb, 'Statistics',
      ['id', 'Label', 'Family', 'Question (current)', 'Proposed question (draft)',
       'Proposed takeaway (N/S example)', 'N/S line', 'N/S vs population', 'W/E line',
       'W/E vs population', 'Themes', 'Days in %d' % YEAR, 'Points', 'Source', 'File licence',
       'Caveats', 'Keep / Change / Drop', 'Your wording', 'Notes'],
      srows, [22, 30, 18, 38, 34, 52, 10, 16, 10, 16, 28, 9, 8, 14, 22, 50, 14, 34, 34],
      edit_cols=(17, 18, 19), keep_col=17)

trows = []
for t in cfg['themes']:
    c = tcount.get(t['id'], {})
    used = sum(1 for d in games['days'] if d['theme'] == t['id'])
    cap = 'filler (no cap)' if t.get('filler') else str(t.get('maxDaysPerYear', cfg['rules']['defaultMaxDaysPerYear']))
    trows.append([t['id'], t['title'], t['blurb'], '; '.join(cat[s]['label'] for s in t['stats']),
                  len(t['stats']), int(c.get('valid_days', 0)), cap, used, '', '', '', ''])
sheet(wb, 'Themes',
      ['id', 'Title', 'Blurb', 'Statistics', 'Stats', 'Valid days', 'Cap / year', 'Days in %d' % YEAR,
       'Keep / Change / Drop', 'New title', 'New blurb', 'Notes'],
      trows, [14, 24, 40, 70, 7, 9, 13, 9, 14, 24, 34, 34], edit_cols=(9, 10, 11, 12), keep_col=9)

crows = []
for d in games['days']:
    dt = datetime.date.fromisoformat(d['date'])
    ax = d['rounds'][0][1]
    crows.append([d['date'], dt.strftime('%a'), (dt - EPOCH).days + 1, d['title'],
                  'North / South' if ax == 'ns' else 'West / East',
                  *[cat[s]['label'] for s, _ in d['rounds']], d['gapKm'],
                  '=HYPERLINK("%s%s","preview")' % (PREVIEW, d['date']), '', ''])
sheet(wb, 'Calendar %d' % YEAR,
      ['Date', 'Day', 'Puzzle #', 'Theme', 'Axis', 'Round 1', 'Round 2', 'Round 3',
       'Closest pair (km)', 'Preview', 'OK?', 'Notes'],
      crows, [11, 6, 9, 22, 13, 32, 32, 32, 10, 10, 8, 34], edit_cols=(11, 12))

out = P('review', 'MapSplit-review.xlsx')
wb.save(out)
print('wrote', os.path.relpath(out, ROOT), '—', len(srows), 'statistics,', len(trows), 'themes,', len(crows), 'days')
